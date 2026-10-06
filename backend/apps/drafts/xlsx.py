"""XLSX export of Compliance Matrix and checklist (specs/09 §6)."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.app_settings import get_setting

from .models import DraftType

XLSX_TYPES = (DraftType.COMPLIANCE_MATRIX, DraftType.BID_CHECKLIST)
VERDICT_COLORS = {"충족": "198754", "보완 필요": "FD7E14", "확인 필요": "0D6EFD"}
STATUS_LABELS = {"TODO": "미완료", "DONE": "완료", "NA": "해당 없음"}
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_thin = Side(style="thin", color="ADB5BD")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
HEAD_FILL = PatternFill("solid", fgColor="E9ECEF")
HIGH_FILL = PatternFill("solid", fgColor="DC3545")
WRAP = Alignment(wrap_text=True, vertical="top")


def _table(ws, start_row: int, columns: list[tuple[str, int]], rows: list[list]) -> int:
    for col, (label, width) in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=col, value=label)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, values in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            cell.alignment = WRAP
    last = start_row + len(rows)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(columns))}{max(last, start_row)}"
    return last


def _footer(ws, row: int) -> None:
    ws.cell(row=row + 2, column=1, value=get_setting("report.disclaimer") or "").font = Font(
        italic=True, color="6C757D", size=9
    )


def compliance_matrix_sheet(wb: Workbook, bid, content: dict) -> None:
    ws = wb.active
    ws.title = "Compliance Matrix"
    header = content.get("header") or {}
    meta = [
        ("공고", f"[{header.get('notice_no') or '-'}] {header.get('bid_title') or bid.title}"),
        ("수요기관", header.get("buyer_org") or "-"),
        (
            "작성",
            f"{header.get('company', '')} {header.get('prepared_by', '')} · {header.get('prepared_at', '')}",
        ),
    ]
    for r, (label, value) in enumerate(meta, start=1):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)
    columns = [
        ("No", 6), ("품목", 22), ("구분", 12), ("요구사항", 40), ("조항", 22),
        ("제안 사양", 30), ("충족 여부", 11), ("응답", 50), ("근거", 30), ("비고", 18),
    ]  # fmt: skip
    rows = content.get("rows") or []
    values = [
        [
            r.get("no"),
            r.get("item"),
            r.get("category"),
            r.get("requirement"),
            r.get("clause"),
            r.get("offered"),
            r.get("compliance"),
            r.get("response"),
            r.get("evidence"),
            r.get("remark"),
        ]  # fmt: skip
        for r in rows
    ]
    start = len(meta) + 2
    last = _table(ws, start, columns, values)
    for i, row in enumerate(rows, start=start + 1):
        verdict = ws.cell(row=i, column=7)
        if row.get("compliance") in VERDICT_COLORS:
            verdict.font = Font(bold=True, color=VERDICT_COLORS[row["compliance"]])
        if row.get("risk") == "HIGH":
            ws.cell(row=i, column=1).fill = HIGH_FILL
            ws.cell(row=i, column=1).font = Font(bold=True, color="FFFFFF")
    _footer(ws, last)


def checklist_sheet(wb: Workbook, bid, content: dict) -> None:
    ws = wb.active
    ws.title = "입찰 체크리스트"
    ws.cell(row=1, column=1, value=f"[{bid.notice_no or '-'}] {bid.title}").font = Font(bold=True)
    columns = [
        ("섹션", 18),
        ("상태", 10),
        ("항목", 60),
        ("기한", 18),
        ("담당", 12),
        ("판정", 11),
        ("비고", 40),
    ]
    values = [
        [
            section.get("title"),
            STATUS_LABELS.get(item.get("status"), item.get("status")),
            item.get("text"),
            item.get("due"),
            item.get("owner"),
            item.get("verdict"),
            item.get("note"),
        ]
        for section in content.get("sections") or []
        for item in section.get("items") or []
    ]
    last = _table(ws, 3, columns, values)
    for r in range(4, last + 1):
        verdict = ws.cell(row=r, column=6)
        if verdict.value in VERDICT_COLORS:
            verdict.font = Font(bold=True, color=VERDICT_COLORS[verdict.value])
        if ws.cell(row=r, column=2).value == "완료":
            ws.cell(row=r, column=2).font = Font(color="198754")
    _footer(ws, last)


def render_xlsx(bid, doc_type: str, content: dict) -> bytes:
    wb = Workbook()
    if doc_type == DraftType.COMPLIANCE_MATRIX:
        compliance_matrix_sheet(wb, bid, content)
    else:
        checklist_sheet(wb, bid, content)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
