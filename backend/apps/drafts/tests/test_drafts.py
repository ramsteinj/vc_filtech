"""Drafts: generators, versions, editing, PDF (specs/09, M7)."""

import re
from urllib.parse import unquote

import pytest

from apps.bids.models import BidItem, BidNotice, BidRequirement
from apps.company.models import Product
from apps.drafts import pdf
from apps.drafts.generators import checklist, compliance_matrix, review_report, technical_query
from apps.drafts.models import DraftDocument, DraftType
from apps.drafts.service import DraftContentError, generate_draft, save_draft, validate_content
from apps.evaluation.service import evaluate_bid
from apps.llm.providers import LLMError

pytestmark = pytest.mark.django_db


@pytest.fixture
def bid(company_data):
    from datetime import datetime

    from django.utils import timezone

    bid = BidNotice.objects.create(
        title="[울산] GT 흡기 필터 구매",
        notice_no="20230323830-00",
        buyer_org="한국동서발전(주)",
        bid_close_at=timezone.make_aware(datetime(2023, 4, 5, 12, 0)),
        qualification_deadline_at=timezone.make_aware(datetime(2023, 4, 4, 18, 0)),
        fit_score=65,
    )
    item = BidItem.objects.create(
        bid=bid,
        item_no="1",
        name="V-bank",
        matched_product=Product.objects.get(model_no="FT-VB500"),
    )
    rows = [
        (item, "DIMENSION", "치수", {"w": 592, "h": 592, "d": 292}),
        (item, "FILTER_GRADE", "ASHRAE 등급", {"standard": "ASHRAE52_2", "class": "MERV 14"}),
        (None, "CERTIFICATION", "ISO 14001", {"cert_type": "ISO14001"}),
        (None, "SUBMISSION_DOC", "하자이행증권", {"doc": "하자이행증권", "timing": "계약 시"}),
        (None, "DELIVERY", "납기", {"days_after_contract": 150}),
        (None, "WARRANTY", "하자보증", {"months": 24}),
    ]
    for order, (it, category, title, norm) in enumerate(rows):
        BidRequirement.objects.create(
            bid=bid,
            item=it,
            category=category,
            title=title,
            requirement_text=title,
            normalized=norm,
            order=order,
        )
    evaluate_bid(bid)
    return bid


def test_compliance_matrix_rows(bid):
    content, used_llm = compliance_matrix(bid)
    assert used_llm is False
    rows = content["rows"]
    assert [r["no"] for r in rows] == list(range(1, 7))
    assert rows[0]["item"].startswith("1") and rows[-1]["item"] == "공통"  # item rows first
    dim = next(r for r in rows if r["category"] == "제품 치수")
    assert dim["compliance"] == "충족" and dim["offered"].startswith("FT-VB500")
    assert "FT-VB500" in dim["evidence"]
    cert = next(r for r in rows if r["category"] == "인증·자격")
    assert cert["risk"] == "HIGH" and cert["compliance"] == "보완 필요"
    assert content["header"]["company"] == "(주)필텍"


def test_checklist_sections(bid):
    content, _ = checklist(bid)
    titles = [s["title"] for s in content["sections"]]
    assert titles[0] == "입찰 참가 자격" and titles[-1] == "주요 일정"
    sections = {s["title"]: s["items"] for s in content["sections"]}
    assert any("ISO 14001" in i["text"] for i in sections["입찰 참가 자격"])
    assert any("하자이행증권" in i["text"] for i in sections["계약 시 제출서류"])
    schedule = [i["text"] for i in sections["주요 일정"]]
    assert schedule[:2] == ["자격(실적) 서류 제출 마감", "전자입찰서 제출 마감"]
    assert sections["주요 일정"][1]["due"] == "2023-04-05 12:00"


def test_technical_query_sources(bid):
    content, _ = technical_query(bid)
    titles = {q["question"] for q in content["questions"]}
    assert any("하자보증" in t for t in titles)  # NEEDS_CONFIRMATION (no rule, no LLM)
    assert len(content["questions"]) == 3  # warranty, delivery + ASHRAE (supplement w/ alternative)
    warranty = next(q for q in content["questions"] if "하자보증" in q["question"])
    assert warranty["background"] == ""  # internal "unjudged" note stays out of the letter
    assert content["header"]["subject"].startswith("[20230323830-00]")


def test_review_report_rule_recommendation(bid):
    content, _ = review_report(bid)
    assert content["recommendation"] == "CONDITIONAL"  # ISO 14001 expired (HIGH)
    assert content["stats"]["high_risk"] == 1
    assert content["fit"]["matched_products"] == ["FT-VB500"]
    assert any("ISO 14001" in r for r in content["key_risks"])


def test_llm_polish_and_failure(bid, fake_llm):
    def respond(request):
        props = request.json_schema["properties"]
        if "rows" in props:
            ids = [int(i) for i in re.findall(r'<row id="(\d+)">', request.user)]
            return {"rows": [{"id": i, "response": f"다듬은 응답 {i}", "remark": ""} for i in ids]}
        raise LLMError("down")

    fake_llm.reset(responder=respond)
    content, used = compliance_matrix(bid)
    assert used and all(r["response"].startswith("다듬은 응답") for r in content["rows"])
    report, used = review_report(bid)  # LLM error → rule content
    assert used is False and report["recommendation"] == "CONDITIONAL"


def test_versions_and_modified_protection(bid, manager_user):
    first = generate_draft(bid, "COMPLIANCE_MATRIX", user=manager_user)
    assert first.version == 1
    again = generate_draft(bid, "COMPLIANCE_MATRIX")
    assert again.pk == first.pk  # overwrite same version
    content = again.content
    content["rows"][0]["remark"] = "담당자 메모"
    save_draft(again, content=content, user=manager_user)
    assert again.is_modified
    regenerated = generate_draft(bid, "COMPLIANCE_MATRIX")
    assert regenerated.version == 2  # edited version kept
    assert DraftDocument.objects.get(pk=first.pk).content["rows"][0]["remark"] == "담당자 메모"
    assert generate_draft(bid, "COMPLIANCE_MATRIX", new_version=True).version == 3


def test_validate_content():
    with pytest.raises(DraftContentError):
        validate_content("COMPLIANCE_MATRIX", {"rows": []})
    with pytest.raises(DraftContentError):
        validate_content(
            "BID_CHECKLIST", {"sections": [{"title": "x", "items": [{"status": "BAD"}]}]}
        )
    with pytest.raises(DraftContentError):
        validate_content("REVIEW_REPORT", {"recommendation": "MAYBE"})
    assert validate_content("TECHNICAL_QUERY", {"header": {}, "questions": []})


def test_pdf_fonts_landscape_and_disposition(bid):
    from weasyprint import HTML

    from apps.drafts.service import build_content
    from apps.drafts.views import evaluation_rows

    contents = {t: build_content(bid, t) for t in DraftType.values}
    assert not DraftDocument.objects.exists()  # build_content never saves
    html = pdf.report_html(bid, contents, evaluation_rows(bid))
    document = HTML(string=html).render()
    widths = {round(p.width) for p in document.pages}
    assert len(widths) == 2  # portrait + landscape (Compliance Matrix)
    fonts = {f for page in document.pages for f in _fonts(page)}
    assert any("Noto Sans KR" in f for f in fonts)
    assert "가상 자료 기반" in html and "입찰 검토 통합 보고서" in html

    header = pdf.content_disposition(pdf.filename(bid, "REPORT"))
    assert header.startswith('attachment; filename="20230323830-00__')
    assert unquote(header.split("UTF-8''")[1]).startswith("20230323830-00_통합보고서_")


def _fonts(page):
    def walk(box):
        style = getattr(box, "style", None)
        if style is not None:
            yield from style["font_family"]
        for child in getattr(box, "children", []) or []:
            yield from walk(child)

    return set(walk(page._page_box))


def test_api_flow(manager_client, bid, inline_jobs):
    base = f"/api/bids/{bid.pk}/drafts"
    listed = manager_client.get(base).data
    assert [d["doc_type"] for d in listed] == [
        "COMPLIANCE_MATRIX",
        "BID_CHECKLIST",
        "TECHNICAL_QUERY",
        "REVIEW_REPORT",
    ]
    assert all(d["latest"] is None for d in listed)
    assert manager_client.get(f"{base}/COMPLIANCE_MATRIX").status_code == 404
    assert manager_client.post(f"{base}/UNKNOWN/generate").status_code == 404

    response = manager_client.post(
        f"{base}/BID_CHECKLIST/generate", {"new_version": False}, format="json"
    )
    assert response.status_code == 202
    assert response.data["job"]["status"] == "SUCCEEDED", response.data["job"]
    draft = manager_client.get(f"{base}/BID_CHECKLIST").data
    assert draft["version"] == 1 and draft["used_llm"] is False

    content = draft["content"]
    content["sections"][0]["items"][0]["status"] = "DONE"
    saved = manager_client.put(
        f"{base}/BID_CHECKLIST", {"content": content, "status": "FINAL"}, format="json"
    )
    assert (
        saved.status_code == 200 and saved.data["is_modified"] and saved.data["status"] == "FINAL"
    )
    bad = manager_client.put(f"{base}/BID_CHECKLIST", {"content": {"sections": "x"}}, format="json")
    assert bad.status_code == 400

    manager_client.post(f"{base}/BID_CHECKLIST/generate", {}, format="json")
    versions = manager_client.get(f"{base}/BID_CHECKLIST/versions").data
    assert [v["version"] for v in versions] == [2, 1]
    old = manager_client.get(f"{base}/BID_CHECKLIST", {"version": 1}).data
    assert old["content"]["sections"][0]["items"][0]["status"] == "DONE"

    response = manager_client.get(f"{base}/BID_CHECKLIST/pdf", {"version": 1})
    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert b"".join(response).startswith(b"%PDF")
    assert "filename*=UTF-8''" in response["Content-Disposition"]

    report = manager_client.get(f"/api/bids/{bid.pk}/report.pdf")
    assert report.status_code == 200 and b"".join(report).startswith(b"%PDF")


def test_api_requires_login(api_client, bid):
    assert api_client.get(f"/api/bids/{bid.pk}/drafts").status_code == 401


def test_empty_llm_letter_is_retried(bid, fake_llm):
    empty = {
        "header": {"to": "", "from": "", "date": "", "subject": ""},
        "intro": "",
        "questions": [],
        "closing": "",
    }
    full = {
        **empty,
        "intro": "질의드립니다.",
        "questions": [
            {
                "no": 1,
                "reference": "",
                "question": "확인 부탁드립니다.",
                "background": "",
                "proposed_alternative": "",
                "requirement_id": 0,
            }
        ],
    }
    fake_llm.reset(empty, full)
    content, used = technical_query(bid)
    assert used and content["questions"][0]["question"] == "확인 부탁드립니다."
    fake_llm.reset(empty, empty)
    content, used = technical_query(bid)
    assert used is False and len(content["questions"]) == 3  # rule-based letter kept


def test_review_report_with_empty_rationale(bid):
    evaluation = bid.requirements.get(category="CERTIFICATION").evaluation
    evaluation.rationale = ""
    evaluation.save()
    content, _ = review_report(bid)
    assert any(r.startswith("ISO 14001: 보완 필요") for r in content["key_risks"])


def _load_xlsx(response):
    from io import BytesIO

    from openpyxl import load_workbook

    return load_workbook(BytesIO(b"".join(response))).active


def test_xlsx_export(manager_client, bid):
    cm = generate_draft(bid, "COMPLIANCE_MATRIX")
    cl = generate_draft(bid, "BID_CHECKLIST")
    base = f"/api/bids/{bid.pk}/drafts"

    response = manager_client.get(f"{base}/COMPLIANCE_MATRIX/xlsx")
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    assert unquote(response["Content-Disposition"].split("UTF-8''")[1]).endswith(".xlsx")
    ws = _load_xlsx(response)
    assert ws.title == "Compliance Matrix"
    header = [c.value for c in ws[5]]
    assert header[:4] == ["No", "품목", "구분", "요구사항"]
    rows = list(ws.iter_rows(min_row=6, max_row=5 + len(cm.content["rows"]), values_only=True))
    assert len(rows) == 6
    high = next(i for i, r in enumerate(cm.content["rows"]) if r["risk"] == "HIGH")
    assert ws.cell(row=6 + high, column=1).fill.fgColor.rgb.endswith("DC3545")
    assert ws.freeze_panes == "A6"
    assert any("가상 자료 기반" in str(c.value) for row in ws.iter_rows() for c in row if c.value)

    content = cl.content
    content["sections"][0]["items"][0]["status"] = "DONE"
    save_draft(cl, content=content)
    ws = _load_xlsx(manager_client.get(f"{base}/BID_CHECKLIST/xlsx"))
    first = [c.value for c in ws[4]]
    assert first[0] == "입찰 참가 자격" and first[1] == "완료"

    assert manager_client.get(f"{base}/TECHNICAL_QUERY/xlsx").status_code == 404
    assert manager_client.get(f"{base}/COMPLIANCE_MATRIX/xlsx", {"version": 9}).status_code == 404
