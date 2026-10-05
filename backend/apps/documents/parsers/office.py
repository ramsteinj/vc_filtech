"""TXT, DOCX, DOC (via LibreOffice) and spreadsheet parsers (specs/06 §2)."""

import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl
import xlrd
from charset_normalizer import from_bytes
from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from .base import (
    ParseError,
    ParseResult,
    UnsupportedFormat,
    normalize_text,
    rows_to_text,
    trim_rows,
)


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    best = from_bytes(data).best()
    if best is None:
        raise ParseError("텍스트 파일의 인코딩을 판별할 수 없습니다.")
    return str(best)


def parse_txt(path: str) -> ParseResult:
    text = normalize_text(decode_text(Path(path).read_bytes()))
    return ParseResult(text=text, pages=[text])


def parse_docx(path: str) -> ParseResult:
    try:
        document = DocxDocument(path)
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise ParseError("DOCX 파일을 열 수 없습니다.") from exc

    lines: list[str] = []
    tables: list[dict] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            lines.append(block.text)
        elif isinstance(block, Table):
            rows = []
            for row in block.rows:
                cells = []
                for cell in row.cells:
                    text = cell.text.strip()
                    if not cells or cells[-1] != text:  # merged cells repeat their text
                        cells.append(text)
                rows.append(cells)
            rows = trim_rows(rows)
            if rows:
                tables.append({"page": None, "rows": rows})
                lines.extend(["[표]", rows_to_text(rows), "[/표]"])
    text = normalize_text("\n".join(lines))
    return ParseResult(text=text, pages=[text], tables=tables)


def parse_doc(path: str) -> ParseResult:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise UnsupportedFormat("구형 .doc 형식은 지원하지 않습니다. DOCX로 저장한 뒤 올려주세요.")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "docx", "--outdir", tmp, path],
            check=True,
            capture_output=True,
            timeout=120,
        )
        converted = next(Path(tmp).glob("*.docx"), None)
        if converted is None:
            raise ParseError(".doc 파일을 변환하지 못했습니다.")
        return parse_docx(str(converted))


def _sheet_result(sheets: list[tuple[str, list[list]]]) -> ParseResult:
    lines: list[str] = []
    tables: list[dict] = []
    for name, raw_rows in sheets:
        rows = trim_rows(raw_rows)
        if not rows:
            continue
        tables.append({"page": None, "sheet": name, "rows": rows})
        lines.extend([f"[시트 {name}]", rows_to_text(rows)])
    text = normalize_text("\n".join(lines))
    return ParseResult(text=text, pages=[text], tables=tables)


def _cell_value(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def parse_xlsx(path: str) -> ParseResult:
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises several unrelated exception types
        raise ParseError("XLSX 파일을 열 수 없습니다.") from exc
    try:
        sheets = [
            (ws.title, [[_cell_value(v) for v in row] for row in ws.iter_rows(values_only=True)])
            for ws in workbook.worksheets
        ]
    finally:
        workbook.close()
    return _sheet_result(sheets)


def parse_xls(path: str) -> ParseResult:
    try:
        workbook = xlrd.open_workbook(path)
    except xlrd.XLRDError as exc:
        raise ParseError("XLS 파일을 열 수 없습니다.") from exc
    sheets = [
        (
            sheet.name,
            [[_cell_value(v) for v in sheet.row_values(r)] for r in range(sheet.nrows)],
        )
        for sheet in workbook.sheets()
    ]
    return _sheet_result(sheets)
