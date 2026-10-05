"""Parser checks against the real initial-data files (specs/12 §5)."""

import re
import zipfile
from pathlib import Path

import pytest
from django.conf import settings
from docx import Document as DocxDocument

from apps.documents.loader import is_ignored
from apps.documents.parsers import ParseError, detect_format, parse_file

DATA = Path(settings.INITIAL_DATA_DIR)
BROKEN = ["捤獥汤捯", "氠瑢", "漠杳", "浵╦", "湰灧"]
SCANNED = {"제한경쟁사유서.pdf", "제한 경쟁 사유서.pdf"}

BID_FILES = sorted(p for p in (DATA / "bid_sample").rglob("*") if p.is_file() and not is_ignored(p))
COMPANY_FILES = sorted((DATA / "company").rglob("*.pdf"))


def test_sample_counts():
    formats = [detect_format(p.name) for p in BID_FILES]
    assert formats.count("HWP") == 24
    assert formats.count("HWPX") == 1
    assert formats.count("XLS") == 2
    assert formats.count("XLSX") == 3
    assert formats.count("PDF") == 2
    assert len(COMPANY_FILES) == 19


@pytest.mark.parametrize("path", BID_FILES, ids=lambda p: p.name)
def test_bid_sample_parses_cleanly(path):
    result = parse_file(str(path), detect_format(path.name))
    if path.name in SCANNED:
        assert result.is_scanned and result.text == ""
        return
    assert not result.is_scanned
    assert len(result.text) > (50 if path.stat().st_size < 20_000 else 100)
    assert not [b for b in BROKEN if b in result.text]
    # Abnormal amounts of CJK ideographs mean control characters were decoded as text.
    hanja = len(re.findall(r"[一-鿿]", result.text))
    assert hanja < max(60, len(result.text) * 0.01)


@pytest.mark.parametrize("path", COMPANY_FILES, ids=lambda p: p.name)
def test_company_pdfs_have_no_watermark_noise(path):
    result = parse_file(str(path), "PDF")
    assert result.tables
    joined = " ".join(cell for t in result.tables for row in t["rows"] for cell in row)
    assert "E가PA" not in joined and "SAMPLE" not in joined


def test_hwp_tables_keep_row_structure():
    path = DATA / "bid_sample/won/20230342721/정비용 기자재 구매규격서.hwp"
    result = parse_file(str(path), "HWP")
    rows = [row for t in result.tables for row in t["rows"]]
    assert ["품번", "품 명", "규 격", "단위", "수량", "자재번호", "비 고"] in rows
    assert "2 | GT Inlet Air Pulse Filter Cartridge | Conical Type" in result.text


def test_spreadsheet_rows():
    path = DATA / "bid_sample/won/20131021333/20131021333-00_1382440245281_구매내역서.xls"
    result = parse_file(str(path), "XLS")
    rows = result.tables[0]["rows"]
    assert rows[1][:3] == ["순번", "품 목", "규 격"]
    assert any("1815" in " ".join(row) for row in rows)


def test_txt_detects_cp949(tmp_path):
    path = tmp_path / "notice.txt"
    path.write_bytes("입찰 공고\n납품기한: 계약 후 150일".encode("cp949"))
    assert parse_file(str(path), "TXT").text == "입찰 공고\n납품기한: 계약 후 150일"


def test_docx_paragraphs_and_tables(tmp_path):
    document = DocxDocument()
    document.add_paragraph("규격서")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "항목", "사양"
    table.cell(1, 0).text, table.cell(1, 1).text = "치수", "592x592x292"
    path = tmp_path / "spec.docx"
    document.save(path)
    result = parse_file(str(path), "DOCX")
    assert result.tables[0]["rows"] == [["항목", "사양"], ["치수", "592x592x292"]]
    assert "치수 | 592x592x292" in result.text


def test_hwpx_sample_parses():
    path = next((DATA / "bid_sample").rglob("*.hwpx"))
    result = parse_file(str(path), "HWPX")
    assert "중소기업" in result.text and result.tables


def test_corrupt_files_raise_parse_error(tmp_path):
    bad = tmp_path / "bad.hwp"
    bad.write_bytes(b"not an ole file")
    with pytest.raises(ParseError):
        parse_file(str(bad), "HWP")
    bad_zip = tmp_path / "bad.hwpx"
    with zipfile.ZipFile(bad_zip, "w") as archive:
        archive.writestr("mimetype", "x")
    with pytest.raises(ParseError):
        parse_file(str(bad_zip), "HWPX")
